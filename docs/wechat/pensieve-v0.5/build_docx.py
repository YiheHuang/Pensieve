from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "Pensieve-公众号推文-可导入版.docx"

INK = "314C58"
DEEP = "183948"
ACCENT = "4D98A8"
MUTED = "6F8993"
PALE = "EEF6F8"
BORDER = "C7DCE2"
WHITE = "FFFFFF"


def set_font(run, name="Microsoft YaHei", size=11.5, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.35

    for style_name, size, color, before, after in [
        ("Heading 1", 18, DEEP, 18, 8),
        ("Heading 2", 14, DEEP, 14, 6),
        ("Heading 3", 12, ACCENT, 10, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def set_cell_shading(paragraph, fill=PALE, left_border=ACCENT):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), left_border)
    borders.append(left)


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, bold=True, color=DEEP)
        rest = p.add_run(text[len(bold_prefix):])
        set_font(rest)
    else:
        set_font(p.add_run(text))
    return p


def add_kicker(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size=9.5, color=ACCENT, bold=True)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(text, style="Heading 1")
    return p


def add_quote(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.keep_together = True
    set_cell_shading(p)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        set_font(p.add_run(line), name="SimSun", size=13.5, color=DEEP, bold=True)
    return p


def add_image(doc, filename, caption, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(ASSETS / filename), width=Inches(6.5))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    cap.paragraph_format.keep_together = True
    set_font(cap.add_run(caption), size=9, color=MUTED)


def add_hyperlink(paragraph, text, url, color=ACCENT, bold=False):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{key}"), "Microsoft YaHei")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)
    configure_styles(doc)

    props = doc.core_properties
    props.title = "我把《哈利·波特》的“冥想盆”，做成了一款私人记忆软件"
    props.subject = "一和札记 · 创造"
    props.author = "一和"
    props.keywords = "Pensieve, 冥想盆, 创造, 记忆"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True
    set_font(title.add_run("我把《哈利·波特》的“冥想盆”，\n做成了一款私人记忆软件"), name="SimSun", size=24, color=DEEP, bold=True)

    byline = doc.add_paragraph()
    byline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    byline.paragraph_format.space_after = Pt(14)
    byline.paragraph_format.keep_with_next = True
    set_font(byline.add_run("一和札记 ·「创造」"), size=10, color=ACCENT, bold=True)

    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.line_spacing = 1.45
    lead.paragraph_format.space_after = Pt(16)
    lead.paragraph_format.keep_with_next = True
    set_font(lead.add_run("保存文字、照片、视频与声音；用 AI 整理情绪和线索；\n通过一句模糊的描述找回往事；再把一整段时间汇聚成报告。"), name="SimSun", size=13.5, color=DEEP)
    add_body(doc, "这就是 Pensieve：一款面向 Windows 的本地优先私人记忆应用。")
    add_image(doc, "home.jpg", "Pensieve · 私人记忆星河", "脱敏后的 Pensieve 主界面")

    add_kicker(doc, "01 · 快速提取")
    add_heading(doc, "只写内容，其余交给 AI")
    add_body(doc, "记录一段记忆时，默认界面只保留一个正文输入区。你可以写下一段话，也可以加入照片、视频、音频，或者直接录下一段声音。")
    add_body(doc, "标题、情绪、时间和标签都放进高级选项，需要时再填写。点击“提取”后，应用会等待 AI 返回完整结果，再进入记忆详情。")
    add_image(doc, "demo-capture.gif", "魔杖携带记忆微光，缓缓浸入水面", "记忆提取与魔杖入水动画")
    add_quote(doc, ["AI 会提取标题、主情绪与副情绪，", "识别人物、地点和主题，并在本地标准化标签。"])
    add_body(doc, "原始正文保持原样。AI 负责整理线索，而不是生成摘要、改写措辞或替你重新讲述经历。")

    add_kicker(doc, "02 · 沉浸寻找")
    add_heading(doc, "记得很模糊，也能把它找回来")
    add_body(doc, "回忆往往只剩下一些场景，例如“下雨时去过的旧书店”“和朋友在湖边散步的那一次”或“最近让我觉得很温暖的事情”。")
    add_body(doc, "在“沉浸”页面，可以直接用这样的自然语言寻找记忆，也可以继续按主情绪或副情绪、起止日期、媒体类型和爱心状态缩小范围。")
    add_image(doc, "demo-search.gif", "沿着文字、情绪与场景，让旧日微光重新浮现", "自然语言沉浸检索动画")
    add_body(doc, "找到记忆后，可以进入沉浸回放。图片会应用低亮、冷色的“回忆滤镜”，与正文分区显示；音频、视频和文字在同一环境中播放。")
    add_body(doc, "正文、发生时间和附件都可以继续增删修改，也可以重新运行 AI 分析。重新分析只更新标题、情绪与标签。")

    add_kicker(doc, "03 · 多元情绪")
    add_heading(doc, "一段记忆可以同时拥有多种情绪")
    add_body(doc, "Pensieve 使用欣喜、宁静、温暖、怀念、勇敢与难过六类标准情绪。AI 会给出一个主情绪，也可以给出若干副情绪。")
    add_body(doc, "这些标签会真正参与搜索、筛选和统计。人物、地点和主题标签也会分别存储并在本地去重。")
    add_image(doc, "feature-emotions.jpg", "一个主情绪，也可以带有若干副情绪", "主情绪、副情绪与标准情绪光谱示意")

    add_kicker(doc, "04 · 时光回响")
    add_heading(doc, "用 AI 总结一段时间")
    add_body(doc, "选择一个北京时间范围，Pensieve 会从记忆数量、活跃天数、附件数量、情绪分布、人物关系、地点场景、主题事件、珍贵片段和变化洞察等维度生成报告。")
    add_image(doc, "demo-time-echoes.gif", "轻触左右方向，让每一页从水下浮现", "时光回响沉浸式分页报告动画")
    add_body(doc, "报告会作为独立快照保存进档案库，可以重命名、收藏或删除。阅读时每次显示一个章节，通过左右按钮、方向键或触屏滑动切换。代表性片段还可以直接“沉入原记忆”。")

    add_kicker(doc, "05 · 长廊与珍藏")
    add_heading(doc, "浏览、珍藏、回收与重新浮现")
    add_body(doc, "记忆长廊提供时间线和网格视图；点亮爱心后，记忆会进入“最珍贵记忆”筛选；主页“今日回响”会从真实的本地记忆中带回一个旧日片段。")
    add_body(doc, "删除先进入回收站，可以恢复；彻底删除会再次确认，并同步清理记录、索引和附件。")
    add_image(doc, "feature-gallery.jpg", "时间线、珍藏筛选与真实记忆卡片", "Pensieve 记忆长廊界面")

    add_kicker(doc, "06 · 附件与备份")
    add_heading(doc, "附件会复制进加密记忆库")
    add_body(doc, "加入照片、音频或视频后，桌面版会把附件复制进应用的加密存储区域，而不是只保存原路径。原文件移动、改名或被清理后，记忆中的附件仍然存在。")
    add_body(doc, "导出 .pensieve 备份时，全部记忆、媒体附件和时光回响档案会一起打包；恢复时也会完整带回。")
    add_image(doc, "feature-backup.jpg", "记忆、附件与时光回响档案一起迁徙", "Pensieve 备份与恢复界面")

    add_kicker(doc, "07 · 本地优先")
    add_heading(doc, "AI 由用户主动启用")
    add_body(doc, "Windows 桌面版使用本地加密数据库保存记忆、报告与附件。AI 默认关闭，可以配置自己的 OpenAI 兼容 API 地址、模型和密钥。")
    add_body(doc, "生成时光回响时，图片、视频和音频原文件留在本地。应用还提供 PIN 锁定、中英文界面、减少动态效果和 Windows 高对比度适配。")
    add_image(doc, "feature-settings.jpg", "语言、守护与 AI 连接都由用户配置", "Pensieve 偏好与守护界面")

    add_kicker(doc, "08 · 使用场景")
    add_heading(doc, "它可以用来做什么？")
    add_body(doc, "Pensieve 可以成为旅行中的照片与声音札记、家庭私人档案、成长记录、灵感与梦境收藏、阶段性情绪回顾，或者一个不必发布到社交网络的个人记忆空间。")
    add_image(doc, "feature-use-cases.jpg", "旅行、家庭、成长、灵感、回顾与私人空间", "Pensieve 使用场景示意")

    add_kicker(doc, "09 · 开始使用")
    add_heading(doc, "现在，它可以被你使用了")
    add_body(doc, "下载并安装 Pensieve，设置自己的称呼与 4—8 位数字 PIN，在“偏好与守护”中按需启用 AI，然后进入“提取”保存第一段记忆。")
    add_image(doc, "feature-start.jpg", "通过 PIN 唤醒自己的私人记忆库", "Pensieve 锁定与唤醒界面")
    links = doc.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links.paragraph_format.space_before = Pt(8)
    links.paragraph_format.space_after = Pt(18)
    add_hyperlink(links, "下载 Pensieve", "https://github.com/YiheHuang/Pensieve/releases/latest", bold=True)
    set_font(links.add_run("　·　"), color=MUTED)
    add_hyperlink(links, "项目主页", "https://github.com/YiheHuang/Pensieve")
    set_font(links.add_run("　·　"), color=MUTED)
    add_hyperlink(links, "使用手册", "https://github.com/YiheHuang/Pensieve/blob/main/docs/%E4%BD%BF%E7%94%A8%E6%89%8B%E5%86%8C.md")
    add_body(doc, "如果你喜欢这种以沉浸体验保存与寻找记忆的方式，欢迎试用 Pensieve，也欢迎在 GitHub 提交建议或参与改进。")

    sig = doc.add_paragraph()
    sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig.paragraph_format.space_before = Pt(18)
    sig.paragraph_format.space_after = Pt(18)
    set_font(sig.add_run("一和 · 写于「一和札记 · 创造」"), size=9.5, color=MUTED)
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(0)
    set_font(note.add_run("说明：Pensieve 是独立开发的开源个人项目，灵感来自“以容器承载和重访记忆”的幻想意象，与 J.K. Rowling、Warner Bros.、Wizarding World 及其关联方无隶属或合作关系。相关作品名与商标归各自权利人所有。"), size=8.5, color=MUTED)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
